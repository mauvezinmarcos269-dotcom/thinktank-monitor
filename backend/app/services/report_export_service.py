from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BODY_FONT = "Microsoft YaHei"
TITLE_FONT = "SimHei"


def _set_run_font(
    run: Any,
    *,
    size: float | None = None,
    bold: bool | None = None,
) -> None:
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    if size is not None:
        run.font.size = Pt(size)

    if bold is not None:
        run.font.bold = bold


def _set_paragraph_spacing(
    paragraph: Any,
    *,
    before: float = 0,
    after: float = 6,
    line_spacing: float = 1.35,
) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing


def _add_page_number(paragraph: Any) -> None:
    paragraph.alignment = 1
    run = paragraph.add_run("第 ")
    _set_run_font(run, size=9)

    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")

    run._r.append(field_begin)
    run._r.append(instr_text)
    run._r.append(field_end)

    suffix = paragraph.add_run(" 页")
    _set_run_font(suffix, size=9)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    for style_name, size in (
        ("Title", 18),
        ("Heading 1", 15),
        ("Heading 2", 13),
    ):
        style = styles[style_name]
        style.font.name = TITLE_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style._element.rPr.rFonts.set(qn("w:eastAsia"), TITLE_FONT)

    footer = section.footer.paragraphs[0]
    _add_page_number(footer)


def normalize_text(value: str | None) -> str:
    if not value:
        return ""

    return value.replace("\r\n", "\n").strip()


def _is_pdf_report(report: Any) -> bool:
    return bool(
        normalize_text(getattr(report, "pdf_url", ""))
        or getattr(report, "page_count", None)
    )


def _build_metadata_items(report: Any) -> list[tuple[str, str]]:
    source_url = normalize_text(getattr(report, "url", ""))
    pdf_url = normalize_text(getattr(report, "pdf_url", ""))
    published_at = getattr(report, "published_at", None)
    ai_generated_at = getattr(report, "ai_generated_at", None)
    document_type = "PDF 报告" if _is_pdf_report(report) else "网页长文"

    metadata = [
        ("原文类型", document_type),
        ("原文链接", source_url or "暂无"),
    ]

    if _is_pdf_report(report):
        metadata.append(("PDF 链接", pdf_url or "暂无"))
    else:
        metadata.append(("PDF", "不适用"))

    metadata.extend(
        [
            ("发布时间", str(published_at or "未知")),
            ("AI 生成时间", str(ai_generated_at or "未知")),
        ]
    )

    return metadata


def build_report_markdown(report: Any) -> str:
    title = normalize_text(getattr(report, "title", "未命名报告"))
    metadata = [
        f"- {label}：{value}"
        for label, value in _build_metadata_items(report)
    ]

    sections = [
        f"# {title}",
        "",
        "## 基本信息",
        *metadata,
        "",
        "## 观点摘要",
        normalize_text(getattr(report, "summary", None)) or "暂无观点摘要。",
        "",
        "## 分析评论",
        normalize_text(getattr(report, "commentary", None)) or "暂无分析评论。",
        "",
        "## 全文翻译",
        normalize_text(getattr(report, "translation", None)) or "暂无全文翻译。",
        "",
        "## 原文正文",
        normalize_text(getattr(report, "content", None)) or "暂无原文正文。",
        "",
    ]

    return "\n".join(sections)


def build_export_filename(
    report_id: int,
    title: str | None,
    extension: str,
) -> str:
    base = title or f"report-{report_id}"
    safe = "".join(
        char
        if char.isalnum() or char in {" ", "-", "_"}
        else "_"
        for char in base
    ).strip()
    safe = "_".join(safe.split())

    if not safe:
        safe = f"report-{report_id}"

    return f"{report_id}-{safe[:80]}-ai-results.{extension}"


def _add_multiline_text(
    document: Document,
    text: str,
    *,
    fallback: str,
) -> None:
    normalized = normalize_text(text)

    if not normalized:
        document.add_paragraph(fallback)
        return

    for block in normalized.split("\n\n"):
        block = block.strip()

        if not block:
            continue

        paragraph = document.add_paragraph()
        _set_paragraph_spacing(paragraph)
        lines = block.split("\n")

        for index, line in enumerate(lines):
            if index > 0:
                paragraph.add_run().add_break(WD_BREAK.LINE)

            paragraph.add_run(line)


def _add_heading(
    document: Document,
    title: str,
    *,
    level: int,
) -> None:
    paragraph = document.add_heading(title, level=level)
    _set_paragraph_spacing(
        paragraph,
        before=10 if level == 1 else 6,
        after=8 if level == 1 else 4,
        line_spacing=1.2,
    )


def _add_static_toc(document: Document) -> None:
    _add_heading(document, "目录", level=1)

    for index, title in enumerate(
        (
            "基本信息",
            "观点摘要",
            "分析评论",
            "全文翻译",
            "原文正文",
        ),
        start=1,
    ):
        paragraph = document.add_paragraph()
        _set_paragraph_spacing(paragraph, after=3)
        run = paragraph.add_run(f"{index}. {title}")
        _set_run_font(run, size=10.5)


def _add_metadata_table(
    document: Document,
    report: Any,
) -> None:
    _add_heading(document, "基本信息", level=1)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"

    metadata = _build_metadata_items(report)

    for label, value in metadata:
        row = table.add_row()
        label_cell, value_cell = row.cells
        label_cell.text = label
        value_cell.text = value

        for paragraph in label_cell.paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, size=9.5, bold=True)

        for paragraph in value_cell.paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, size=9.5)


def _add_section(
    document: Document,
    title: str,
    text: str | None,
    *,
    fallback: str,
    page_break_before: bool = False,
) -> None:
    if page_break_before:
        document.add_section(WD_SECTION.NEW_PAGE)

    _add_heading(document, title, level=1)
    _add_multiline_text(
        document,
        text or "",
        fallback=fallback,
    )


def build_report_docx(report: Any) -> bytes:
    document = Document()
    _configure_document(document)

    title = normalize_text(getattr(report, "title", None)) or "未命名报告"

    title_paragraph = document.add_heading(title, level=0)
    title_paragraph.alignment = 1
    _set_paragraph_spacing(
        title_paragraph,
        before=24,
        after=18,
        line_spacing=1.2,
    )

    subtitle = document.add_paragraph()
    subtitle.alignment = 1
    subtitle_run = subtitle.add_run("智库报告翻译与分析成果")
    _set_run_font(subtitle_run, size=12, bold=True)
    _set_paragraph_spacing(subtitle, after=18, line_spacing=1.2)

    _add_static_toc(document)
    document.add_section(WD_SECTION.NEW_PAGE)

    _add_metadata_table(document, report)

    _add_section(
        document,
        "观点摘要",
        getattr(report, "summary", None),
        fallback="暂无观点摘要。",
    )

    _add_section(
        document,
        "分析评论",
        getattr(report, "commentary", None),
        fallback="暂无分析评论。",
    )

    _add_section(
        document,
        "全文翻译",
        getattr(report, "translation", None),
        fallback="暂无全文翻译。",
        page_break_before=True,
    )

    _add_section(
        document,
        "原文正文",
        getattr(report, "content", None),
        fallback="暂无原文正文。",
        page_break_before=True,
    )

    output = BytesIO()
    document.save(output)

    return output.getvalue()
