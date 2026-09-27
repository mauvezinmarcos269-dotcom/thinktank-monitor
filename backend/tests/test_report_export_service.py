from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document

from app.api.v1 import reports as reports_api
from app.schemas.report import ReportBatchExportRequest
from app.services.report_export_service import (
    build_export_filename,
    build_report_docx,
    build_report_markdown,
    normalize_text,
)


def _sample_report():
    return SimpleNamespace(
        id=1,
        title="Sample China Report",
        url="https://example.com/report",
        pdf_url="https://example.com/report.pdf",
        content_kind="pdf",
        review_status="approved",
        review_note="已复核，可采用。",
        reviewed_at=datetime(2026, 1, 4, 5, 6, 7),
        published_at=datetime(2026, 1, 2, 3, 4, 5),
        ai_generated_at=datetime(2026, 1, 3, 4, 5, 6),
        summary="观点摘要正文。",
        commentary="分析评论正文。",
        translation="全文翻译正文。",
        content="Original content.",
    )


def _sample_web_article():
    return SimpleNamespace(
        id=2,
        title="Sample Brookings Web Article",
        url="https://www.brookings.edu/articles/china-policy/",
        pdf_url=None,
        content_kind="web_article",
        review_status="pending_review",
        review_note=None,
        reviewed_at=None,
        page_count=None,
        published_at=datetime(2026, 9, 9, 10, 0, 0),
        ai_generated_at=datetime(2026, 9, 10, 11, 0, 0),
        summary="观点摘要正文。",
        commentary="分析评论正文。",
        translation="全文翻译正文。",
        content="Original web article content.",
    )


def test_build_report_markdown_contains_all_sections():
    markdown = build_report_markdown(_sample_report())

    assert "# Sample China Report" in markdown
    assert "## 基本信息" in markdown
    assert "- 报告 ID：1" in markdown
    assert "- 原文类型：PDF 报告" in markdown
    assert "- 复核状态：已通过" in markdown
    assert "- 复核时间：2026-01-04 05:06:07" in markdown
    assert "- 复核意见：已复核，可采用。" in markdown
    assert "- PDF 链接：https://example.com/report.pdf" in markdown
    assert "- 原文字数：17 字符" in markdown
    assert "- 翻译字数：7 字符" in markdown
    assert "- 主要观点字数：7 字符" in markdown
    assert "- 深层研判字数：7 字符" in markdown
    assert "## 全文翻译稿" in markdown
    assert "## 分析评论稿" in markdown
    assert "### 第一部分：主要观点" in markdown
    assert "### 第二部分：深层研判" in markdown
    assert "## 原文正文" in markdown


def test_build_report_markdown_uses_delivery_section_order():
    markdown = build_report_markdown(_sample_report())

    expected_order = [
        "## 基本信息",
        "## 全文翻译稿",
        "## 分析评论稿",
        "### 第一部分：主要观点",
        "### 第二部分：深层研判",
        "## 原文正文",
    ]

    positions = [
        markdown.index(section)
        for section in expected_order
    ]

    assert positions == sorted(positions)


def test_build_report_markdown_marks_web_article_metadata():
    markdown = build_report_markdown(_sample_web_article())

    assert "- 原文类型：网页长文" in markdown
    assert "- 复核状态：待复核" in markdown
    assert "- 原文链接：https://www.brookings.edu/articles/china-policy/" in markdown
    assert "- PDF：不适用" in markdown
    assert "PDF 链接" not in markdown


def test_build_report_markdown_prefers_content_kind_for_document_type():
    report = _sample_report()
    report.pdf_url = None
    report.page_count = None
    report.content_kind = "pdf"

    markdown = build_report_markdown(report)

    assert "- 原文类型：PDF 报告" in markdown
    assert "- PDF 链接：暂无" in markdown


def test_build_report_docx_contains_formal_sections():
    docx_bytes = build_report_docx(_sample_report())
    document = Document(BytesIO(docx_bytes))
    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    assert paragraphs[0] == "Sample China Report"
    assert "全文翻译稿与分析评论稿" in paragraphs
    assert "目录" in paragraphs
    assert "全文翻译稿" in paragraphs
    assert "分析评论稿" in paragraphs
    assert "分析评论稿第一部分：主要观点" in paragraphs
    assert "分析评论稿第二部分：深层研判" in paragraphs
    assert "原文正文" in paragraphs
    assert len(document.tables) == 1


def test_build_report_docx_marks_web_article_metadata():
    docx_bytes = build_report_docx(_sample_web_article())
    document = Document(BytesIO(docx_bytes))
    table = document.tables[0]
    rows = [(row.cells[0].text, row.cells[1].text) for row in table.rows]

    assert ("原文类型", "网页长文") in rows
    assert ("复核状态", "待复核") in rows
    assert ("PDF", "不适用") in rows
    assert ("原文字数", "29 字符") in rows
    assert ("翻译字数", "7 字符") in rows
    assert not any(label == "PDF 链接" for label, _ in rows)


def test_build_export_filename_sanitizes_title_and_limits_length():
    filename = build_export_filename(
        42,
        "  China: Tech/Industrial Strategy? * 2026 " * 4,
        "docx",
    )

    assert filename.startswith("42-China_")
    assert filename.endswith("-ai-results.docx")
    assert ":" not in filename
    assert "/" not in filename
    assert "?" not in filename
    assert "*" not in filename
    assert len(filename) <= len("42--ai-results.docx") + 80


def test_build_export_filename_falls_back_when_title_is_blank():
    assert (
        build_export_filename(7, " \t\n ", "md")
        == "7-report-7-ai-results.md"
    )


def test_normalize_text_strips_outer_whitespace_and_crlf():
    assert normalize_text("  第一行\r\n第二行  ") == "第一行\n第二行"
    assert normalize_text(None) == ""


def test_build_report_markdown_uses_fallbacks_for_empty_ai_fields():
    report = _sample_report()
    report.title = "   "
    report.summary = None
    report.commentary = ""
    report.translation = "  "
    report.content = None

    markdown = build_report_markdown(report)

    assert "# 未命名报告" in markdown
    assert "暂无主要观点，等待 AI 处理完成。" in markdown
    assert "暂无深层研判，等待 AI 处理完成。" in markdown
    assert "暂无全文翻译。" in markdown
    assert "暂无原文正文。" in markdown
    assert "- 原文字数：暂无" in markdown
    assert "- 翻译字数：暂无" in markdown


def test_build_report_markdown_infers_unknown_pdf_like_content_kind():
    report = _sample_report()
    report.content_kind = "unknown"
    report.pdf_url = "https://example.com/file.pdf"
    report.page_count = None

    markdown = build_report_markdown(report)

    assert "- 原文类型：PDF 报告" in markdown
    assert "- PDF 链接：https://example.com/file.pdf" in markdown


def test_build_report_markdown_infers_unknown_web_like_content_kind():
    report = _sample_web_article()
    report.content_kind = "unknown"
    report.pdf_url = None
    report.page_count = None

    markdown = build_report_markdown(report)

    assert "- 原文类型：网页长文" in markdown
    assert "- PDF：不适用" in markdown


@pytest.mark.asyncio
async def test_batch_export_report_results_builds_markdown_zip(monkeypatch):
    reports = {
        1: _sample_report(),
        2: _sample_web_article(),
    }

    async def fake_get_report(_db, report_id):
        return reports.get(report_id)

    monkeypatch.setattr(reports_api.report_service, "get_report", fake_get_report)

    response = await reports_api.batch_export_report_results(
        ReportBatchExportRequest(
            report_ids=[1, 2, 2, 404],
            export_format="markdown",
        ),
        SimpleNamespace(),
        SimpleNamespace(),
    )

    assert response.media_type == "application/zip"
    assert "thinktank-reports-markdown" in response.headers["Content-Disposition"]

    with ZipFile(BytesIO(response.body)) as archive:
        names = archive.namelist()

        assert len([name for name in names if name.endswith(".md")]) == 2
        assert "missing-report-ids.txt" in names
        assert archive.read("missing-report-ids.txt").decode("utf-8") == "404"
        assert any(name.startswith("1-Sample_China_Report") for name in names)
        assert any(
            name.startswith("2-Sample_Brookings_Web_Article")
            for name in names
        )

        first_markdown = archive.read(
            next(name for name in names if name.startswith("1-"))
        ).decode("utf-8")

    assert "# Sample China Report" in first_markdown
    assert "## 全文翻译稿" in first_markdown


@pytest.mark.asyncio
async def test_batch_export_report_results_builds_docx_zip(monkeypatch):
    async def fake_get_report(_db, report_id):
        return _sample_report() if report_id == 1 else None

    monkeypatch.setattr(reports_api.report_service, "get_report", fake_get_report)

    response = await reports_api.batch_export_report_results(
        ReportBatchExportRequest(report_ids=[1], export_format="docx"),
        SimpleNamespace(),
        SimpleNamespace(),
    )

    with ZipFile(BytesIO(response.body)) as archive:
        names = archive.namelist()

        assert len(names) == 1
        assert names[0].endswith(".docx")
        assert archive.read(names[0]).startswith(b"PK")
